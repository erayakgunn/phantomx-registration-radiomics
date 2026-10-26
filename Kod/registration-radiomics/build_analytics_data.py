import json
import sys
from pathlib import Path

import numpy as np
import openpyxl


def load_report(path: Path, label: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {"label": label, "summary": [], "series": [], "dice": {}, "date_summary": [], "hist": {}, "diagnostics": []}

    ws = wb["Summary"]
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] is None:
            continue
        result["summary"].append({
            "organ": row[0], "metric": row[1], "comparison": row[2], "n": row[3],
            "mean": row[4], "std": row[5], "median": row[6], "min": row[7], "max": row[8],
        })
    for line in ws.iter_rows(min_row=2, max_row=2, values_only=True):
        if line[0]:
            result["meta"] = str(line[0])

    for organ in ("Aorta", "Spleen"):
        sheet = wb[f"{organ}_Dice_Series"]
        rows = list(sheet.iter_rows(min_row=3, values_only=True))
        series = [c for c in rows[0][1:] if c]
        matrix = []
        for row in rows[1:]:
            vals = []
            for v in row[1:]:
                vals.append(None if v is None else round(float(v), 4))
            if any(x is not None for x in vals):
                matrix.append(vals)
        if not result["series"]:
            result["series"] = series
        result["dice"][organ.lower()] = matrix

    ws = wb["Date_Summary"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        result["date_summary"].append({
            "a": row[0], "b": row[1], "n": row[2],
            "aorta_before": row[3], "aorta_after": row[4], "aorta_gain": row[5],
            "spleen_before": row[6], "spleen_after": row[7], "spleen_gain": row[8],
            "aorta_improved": row[9], "spleen_improved": row[10],
        })

    ws = wb["Pairs"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    idx = {name: i for i, name in enumerate(header)}
    a_vals, s_vals = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["Aorta Dice"]] is not None:
            a_vals.append(float(row[idx["Aorta Dice"]]))
        if row[idx["Spleen Dice"]] is not None:
            s_vals.append(float(row[idx["Spleen Dice"]]))
    for organ, vals in (("aorta", a_vals), ("spleen", s_vals)):
        counts, edges = np.histogram(vals, bins=20)
        result["hist"][organ] = {
            "edges": [round(float(e), 4) for e in edges],
            "counts": [int(c) for c in counts],
        }
    result["n_pairs"] = len(a_vals)

    ws = wb["Dice_Diagnostics"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    idx = {name: i for i, name in enumerate(header)}
    diag = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        a = row[idx["Aorta Dice"]]
        s = row[idx["Spleen Dice"]]
        if a is None or s is None:
            continue
        diag.append({
            "a": row[idx["Series A"]], "b": row[idx["Series B"]],
            "same": row[idx["Same session"]],
            "aorta": round(float(a), 4), "spleen": round(float(s), 4),
            "aorta_c": row[idx["Aorta centroid diff (mm)"]],
            "spleen_c": row[idx["Spleen centroid diff (mm)"]],
            "category": row[idx["Category"]], "severity": row[idx["Severity"]],
        })
    diag.sort(key=lambda d: min(d["aorta"], d["spleen"]))
    result["diagnostics"] = diag[:30]
    wb.close()
    return result


def main():
    sep_path, pcct_path, out_path = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
    data = {
        "sep": load_report(sep_path, "SEP"),
        "pcct": load_report(pcct_path, "PCCT"),
    }
    text = "window.ANALYTICS = " + json.dumps(data, ensure_ascii=False) + ";\n"
    out_path.write_text(text, encoding="utf-8")
    for key, ds in data.items():
        print(f"{key}: series={len(ds['series'])} pairs={ds['n_pairs']} date_pairs={len(ds['date_summary'])} "
              f"summary_rows={len(ds['summary'])} diag={len(ds['diagnostics'])}")
    print(f"Wrote {out_path} ({out_path.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
