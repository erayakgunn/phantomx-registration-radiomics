"""Corrected PCCT before/after-registration comparison.

Fixes vs. the student's evaluator:
  * organ ids resolved from the *source* label_map.json (140 kVp: spleen=3, aorta=6),
  * take order taken from MOVE_LOG (take1..take10), identical for before/after,
  * protocol name taken from MOVE_LOG, identical for before/after,
  * baseline take1 is the unregistered reference on both sides.
"""

import json
from collections import defaultdict
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import SimpleITK as sitk
from scipy.ndimage import binary_erosion
from scipy.spatial import cKDTree

SRC = Path(r"C:/Users/eray/Desktop/zurich_pcct")
REG = Path(r"C:/Users/eray/Desktop/zurich_pcct_work/REGISTRATION")
OUT = Path(r"C:/Users/eray/Desktop/KARSILASTIRMA_PCCT_DUZELTILMIS.xlsx")
SESSIONS = ["PCTT_20260929_1", "PCTT_20260929_2"]
PAD = 20


def load_movelog(session):
    data = json.loads((SRC / session / "dicoms" / "MOVE_LOG.json").read_text(encoding="utf-8"))
    m = defaultdict(dict)
    for it in data:
        old = Path(it["eski"])
        new = Path(it["yeni"])
        m[old.parent.parent.name][old.parent.name] = new.parent.name
    return m


def resolve_ids(lm_path):
    lm = json.loads(Path(lm_path).read_text(encoding="utf-8"))
    ids = {}
    for k, v in lm.items():
        if isinstance(v, str):
            if v.lower() == "spleen":
                ids["spleen"] = int(k)
            elif v.lower() == "aorta":
                ids["aorta"] = int(k)
    return ids


def bbox1(mask):
    st = sitk.LabelShapeStatisticsImageFilter()
    st.Execute(mask)
    return st.GetBoundingBox(1) if st.HasLabel(1) else None


def crop(mask, bb, pad=PAD):
    sz = mask.GetSize()
    x0 = max(0, bb[0] - pad); y0 = max(0, bb[1] - pad); z0 = max(0, bb[2] - pad)
    x1 = min(sz[0], bb[0] + bb[3] + pad); y1 = min(sz[1], bb[1] + bb[4] + pad); z1 = min(sz[2], bb[2] + bb[5] + pad)
    return mask[x0:x1, y0:y1, z0:z1]


def surface_points(mask):
    arr = sitk.GetArrayFromImage(mask)
    if not arr.any():
        return np.empty((0, 3), np.float32)
    zs, ys, xs = np.nonzero(arr)
    m = 1
    zmin, zmax = max(0, zs.min() - m), min(arr.shape[0], zs.max() + m + 1)
    ymin, ymax = max(0, ys.min() - m), min(arr.shape[1], ys.max() + m + 1)
    xmin, xmax = max(0, xs.min() - m), min(arr.shape[2], xs.max() + m + 1)
    sub = arr[zmin:zmax, ymin:ymax, xmin:xmax]
    eroded = binary_erosion(sub > 0)
    surface = (sub > 0) & (~eroded)
    zyx = np.argwhere(surface)
    if len(zyx) == 0:
        return np.empty((0, 3), np.float32)
    zyx = zyx + np.array([zmin, ymin, xmin])
    xyz = zyx[:, ::-1]
    origin = np.array(mask.GetOrigin(), np.float32)
    spacing = np.array(mask.GetSpacing(), np.float32)
    direction = np.array(mask.GetDirection(), np.float32).reshape(3, 3)
    return (origin + (xyz.astype(np.float32) * spacing) @ direction.T).astype(np.float32)


def load_item(seg_path, ids):
    im = sitk.ReadImage(str(seg_path))
    s = sitk.Cast(im == ids["spleen"], sitk.sitkUInt8)
    a = sitk.Cast(im == ids["aorta"], sitk.sitkUInt8)
    sp = im.GetSpacing()
    vox = sp[0] * sp[1] * sp[2] / 1000.0
    bs, ba = bbox1(s), bbox1(a)
    s_pts = surface_points(s)
    a_pts = surface_points(a)
    return {
        "sub_s": crop(s, bs) if bs else None,
        "sub_a": crop(a, ba) if ba else None,
        "s_pts": s_pts, "a_pts": a_pts,
        "s_tree": cKDTree(s_pts) if len(s_pts) else None,
        "a_tree": cKDTree(a_pts) if len(a_pts) else None,
        "s_vol": float(np.count_nonzero(sitk.GetArrayFromImage(s)) * vox),
        "a_vol": float(np.count_nonzero(sitk.GetArrayFromImage(a)) * vox),
    }


def distances(pts_a, pts_b, tree_a, tree_b):
    if len(pts_a) == 0 or len(pts_b) == 0 or tree_a is None or tree_b is None:
        return 0.0, 0.0
    dab, _ = tree_b.query(pts_a)
    dba, _ = tree_a.query(pts_b)
    return float(max(np.percentile(dab, 95), np.percentile(dba, 95))), float(0.5 * (dab.mean() + dba.mean()))


def dice_masks(sub_a, sub_b):
    if sub_a is None or sub_b is None:
        return 0.0
    r = sitk.ResampleImageFilter()
    r.SetReferenceImage(sub_a)
    r.SetInterpolator(sitk.sitkNearestNeighbor)
    r.SetDefaultPixelValue(0)
    r.SetOutputPixelType(sitk.sitkUInt8)
    b_res = r.Execute(sub_b)
    A = sitk.GetArrayViewFromImage(sub_a) > 0
    B = sitk.GetArrayViewFromImage(b_res) > 0
    tot = int(A.sum()) + int(B.sum())
    return float(2.0 * np.logical_and(A, B).sum() / tot) if tot else 0.0


def compute_pair(a, b):
    s_hd, s_as = distances(a["s_pts"], b["s_pts"], a["s_tree"], b["s_tree"])
    a_hd, a_as = distances(a["a_pts"], b["a_pts"], a["a_tree"], b["a_tree"])
    return {
        "spleen_dice": round(dice_masks(a["sub_s"], b["sub_s"]), 4),
        "aorta_dice": round(dice_masks(a["sub_a"], b["sub_a"]), 4),
        "spleen_hd95": round(s_hd, 3), "aorta_hd95": round(a_hd, 3),
        "spleen_assd": round(s_as, 3), "aorta_assd": round(a_as, 3),
    }


def build():
    groups = []
    for g in sorted(REG.iterdir()):
        mf = g / "batch_manifest.json"
        if mf.exists():
            groups.append((g, json.loads(mf.read_text(encoding="utf-8"))))

    rows = []
    summary = []
    labels_used = []

    for session in SESSIONS:
        movelog = load_movelog(session)
        src_seg = SRC / session / "segmentation"
        for prot, takes in sorted(movelog.items()):
            fixed = takes["take1"]
            group = next(((g, m) for g, m in groups if m.get("fixed_series") == fixed), None)
            if group is None:
                print(f"  [ATLA] kayit grubu yok: {prot}")
                continue
            gdir, manifest = group
            ids = resolve_ids(src_seg / fixed / "label_map.json")
            labels_used.append((prot, ids["spleen"], ids["aorta"]))

            take_order = sorted(takes, key=lambda x: int(x.replace("take", "")))
            before, after = {}, {}
            for tk in take_order:
                seg = src_seg / takes[tk] / "seg.nii.gz"
                if tk == "take1":
                    before[tk] = load_item(seg, ids)
                    after[tk] = before[tk]
                else:
                    before[tk] = load_item(seg, ids)
                    reg_seg = gdir / f"{gdir.name}__{takes[tk]}" / "seg_registered.nii.gz"
                    after[tk] = load_item(reg_seg, ids)

            for i, ta in enumerate(take_order):
                for tb in take_order[i + 1:]:
                    rb = compute_pair(before[ta], before[tb])
                    ra = compute_pair(after[ta], after[tb])
                    rows.append({
                        "Protokol": prot, "Cift": f"{ta}-{tb}",
                        "Dalak_Once": rb["spleen_dice"], "Dalak_Sonra": ra["spleen_dice"],
                        "Dalak_Fark": round(ra["spleen_dice"] - rb["spleen_dice"], 4),
                        "Aort_Once": rb["aorta_dice"], "Aort_Sonra": ra["aorta_dice"],
                        "Aort_Fark": round(ra["aorta_dice"] - rb["aorta_dice"], 4),
                        "Dalak_HD95_Once": rb["spleen_hd95"], "Dalak_HD95_Sonra": ra["spleen_hd95"],
                        "Aort_HD95_Once": rb["aorta_hd95"], "Aort_HD95_Sonra": ra["aorta_hd95"],
                    })
            sb = [r["Dalak_Once"] for r in rows if r["Protokol"] == prot]
            sa = [r["Dalak_Sonra"] for r in rows if r["Protokol"] == prot]
            ab = [r["Aort_Once"] for r in rows if r["Protokol"] == prot]
            aa = [r["Aort_Sonra"] for r in rows if r["Protokol"] == prot]
            shb = [r["Dalak_HD95_Once"] for r in rows if r["Protokol"] == prot]
            sha = [r["Dalak_HD95_Sonra"] for r in rows if r["Protokol"] == prot]
            ahb = [r["Aort_HD95_Once"] for r in rows if r["Protokol"] == prot]
            aha = [r["Aort_HD95_Sonra"] for r in rows if r["Protokol"] == prot]
            summary.append({
                "Protokol": prot, "Cift": len(sb),
                "Dalak_Once": np.median(sb), "Dalak_Sonra": np.median(sa),
                "Dalak_Fark": np.median(sa) - np.median(sb),
                "Aort_Once": np.median(ab), "Aort_Sonra": np.median(aa),
                "Aort_Fark": np.median(aa) - np.median(ab),
                "Dalak_HD95_Once": np.median(shb), "Dalak_HD95_Sonra": np.median(sha),
                "Aort_HD95_Once": np.median(ahb), "Aort_HD95_Sonra": np.median(aha),
            })
            print(f"  {prot}: dalak {np.median(sb):.4f} -> {np.median(sa):.4f} "
                  f"({np.median(sa)-np.median(sb):+.4f}) | aort {np.median(ab):.4f} -> {np.median(aa):.4f} "
                  f"({np.median(aa)-np.median(ab):+.4f})")
    return rows, summary, labels_used


def write_xlsx(rows, summary, labels_used):
    wb = openpyxl.Workbook()
    navy = PatternFill("solid", start_color="1F4E78")
    gray = PatternFill("solid", start_color="D9E1F2")

    def head(ws, r, cols):
        for c, h in enumerate(cols, 1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF", size=9)
            cell.fill = navy
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws = wb.active
    ws.title = "Ozet"
    ws["A1"] = "PCCT DUZELTILMIS KARSILASTIRMA: KAYITSIZ (ONCE) vs REGISTERED (SONRA)"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    cols = ["Protokol", "Cift", "Dalak Once", "Dalak Sonra", "Dalak Fark",
            "Aort Once", "Aort Sonra", "Aort Fark",
            "Dalak HD95 Once", "Dalak HD95 Sonra", "Aort HD95 Once", "Aort HD95 Sonra"]
    head(ws, 3, cols)
    for r, s in enumerate(summary, 4):
        vals = [s["Protokol"], s["Cift"], s["Dalak_Once"], s["Dalak_Sonra"], s["Dalak_Fark"],
                s["Aort_Once"], s["Aort_Sonra"], s["Aort_Fark"],
                s["Dalak_HD95_Once"], s["Dalak_HD95_Sonra"], s["Aort_HD95_Once"], s["Aort_HD95_Sonra"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(size=9)
            cell.alignment = Alignment(horizontal="center" if c > 2 else "left")
            if c == 5 or c == 8:
                cell.fill = PatternFill("solid", start_color="E2EFDA" if v >= 0 else "FCE4D6")
    ws.column_dimensions["A"].width = 46
    for c in range(2, 13):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.freeze_panes = "C4"

    ws2 = wb.create_sheet("Ciftler")
    ws2["A1"] = f"TUM CIFTLER ({len(rows)})"
    ws2["A1"].font = Font(bold=True, size=12, color="1F4E78")
    cols2 = ["Protokol", "Cift", "Dalak Once", "Dalak Sonra", "Dalak Fark",
             "Aort Once", "Aort Sonra", "Aort Fark",
             "Dalak HD95 Once", "Dalak HD95 Sonra", "Aort HD95 Once", "Aort HD95 Sonra"]
    head(ws2, 3, cols2)
    for r, d in enumerate(rows, 4):
        vals = [d["Protokol"], d["Cift"], d["Dalak_Once"], d["Dalak_Sonra"], d["Dalak_Fark"],
                d["Aort_Once"], d["Aort_Sonra"], d["Aort_Fark"],
                d["Dalak_HD95_Once"], d["Dalak_HD95_Sonra"], d["Aort_HD95_Once"], d["Aort_HD95_Sonra"]]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.font = Font(size=9)
            cell.alignment = Alignment(horizontal="center" if c > 2 else "left")
            if c in (5, 8):
                cell.fill = PatternFill("solid", start_color="E2EFDA" if v >= 0 else "FCE4D6")
    ws2.column_dimensions["A"].width = 46
    for c in range(2, 13):
        ws2.column_dimensions[get_column_letter(c)].width = 14
    ws2.freeze_panes = "C4"

    ws3 = wb.create_sheet("Etiketler")
    ws3["A1"] = "PROTOKOL BASINA KULLANILAN ORGAN ETIKETLERI (kaynak label_map.json)"
    ws3["A1"].font = Font(bold=True, size=12, color="1F4E78")
    head(ws3, 3, ["Protokol", "Spleen ID", "Aorta ID"])
    for r, (p, s, a) in enumerate(labels_used, 4):
        ws3.cell(row=r, column=1, value=p).font = Font(size=9)
        ws3.cell(row=r, column=2, value=s).font = Font(size=9)
        ws3.cell(row=r, column=3, value=a).font = Font(size=9)
    ws3.column_dimensions["A"].width = 46

    wb.save(OUT)
    print(f"\n[OK] {OUT}")


if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--registration", type=Path, default=REG)
    ap.add_argument("--output", type=Path, default=OUT)
    args = ap.parse_args()
    REG = args.registration
    OUT = args.output
    print(f"Hesaplaniyor: {REG}")
    rows, summary, labels_used = build()
    write_xlsx(rows, summary, labels_used)
    print("\n=== PROTOKOL OZETI ===")
    for s in summary:
        print(f"{s['Protokol']}: dalak {s['Dalak_Once']:.4f}->{s['Dalak_Sonra']:.4f} "
              f"({s['Dalak_Fark']:+.4f}) | aort {s['Aort_Once']:.4f}->{s['Aort_Sonra']:.4f} "
              f"({s['Aort_Fark']:+.4f})")
