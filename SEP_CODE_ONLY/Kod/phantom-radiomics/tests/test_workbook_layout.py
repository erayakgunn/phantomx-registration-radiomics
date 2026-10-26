"""Çalışma kitabının sütun düzeni ve beyan sayfası.

Bu testlerin varlık sebebi somut: `_COMPACT_*` sabitleri bir düzenleme sırasında
kazara silindiğinde pipeline sessizce çökmedi — Excel yazımı yarıda kaldı ve
çalışma kitabı öznitelik sayfası olmadan diske yazıldı. Sütun düzeni bir
görünüm ayrıntısı değil, çıktının sözleşmesi.
"""

from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from conftest import PROJECT_DIR, organ_radiomics_project


# Sol meta bloğu, sorduğu soruya göre gruplanmış: hangi ROI / kayıt / çekim.
EXPECTED_ROI_BLOCK = ["label", "insert_size"]
EXPECTED_BOOKKEEPING_BLOCK = ["series", "n_vox"]
EXPECTED_ACQUISITION_BLOCK = [
    "recon_k", "k_strength", "kVp", "mAs", "FOV", "SliceThickness", "me_keV",
    # Test-retest index. The phantom is scanned repeatedly, so the same protocol
    # appears several times and the repeat number is an acquisition condition.
    "repeat",
]
EXPECTED_META_BLOCK = (EXPECTED_ROI_BLOCK + EXPECTED_BOOKKEEPING_BLOCK
                       + EXPECTED_ACQUISITION_BLOCK)
EXPECTED_TRAILING = ["section"]


def test_compact_constants_are_defined(pipeline):
    """Hepsi tanımlı olmalı — biri eksikse _compact_wide NameError veriyor."""
    for name in ("_COMPACT_DROP", "_COMPACT_RENAME",
                 "_COMPACT_META_ORDER", "_COMPACT_TRAILING",
                 "_COMPACT_ROI_COLUMNS", "_COMPACT_BOOKKEEPING_COLUMNS",
                 "_COMPACT_ACQUISITION_COLUMNS"):
        assert hasattr(pipeline, name), f"{name} tanımlı değil"

    assert pipeline._COMPACT_ROI_COLUMNS == EXPECTED_ROI_BLOCK
    assert pipeline._COMPACT_BOOKKEEPING_COLUMNS == EXPECTED_BOOKKEEPING_BLOCK
    assert pipeline._COMPACT_ACQUISITION_COLUMNS == EXPECTED_ACQUISITION_BLOCK
    assert pipeline._COMPACT_META_ORDER == EXPECTED_META_BLOCK
    assert pipeline._COMPACT_TRAILING == EXPECTED_TRAILING


def test_roi_name_leads_the_sheet(pipeline):
    """Tabloyu tararken aranan ilk şey hangi insert/star/texture olduğu."""
    assert pipeline._COMPACT_META_ORDER[0] == "label"
    # Boyut adla birlikte: aynı insert her çapta bir kez görünüyor, yani ikisi
    # olmadan satır tekil değil.
    assert pipeline._COMPACT_META_ORDER[1] == "insert_size"


def test_only_identifying_columns_are_frozen(pipeline):
    """Çekim koşulları da solda ve renkli, ama dondurulmuyor."""
    wide = pd.DataFrame(columns=EXPECTED_META_BLOCK + ["stat_mean", "section"])
    assert pipeline._sheet_meta_count(wide) == len(
        EXPECTED_ROI_BLOCK + EXPECTED_BOOKKEEPING_BLOCK
    )


def test_rename_map_covers_every_renamed_meta_column(pipeline):
    """Görünen meta bloğundaki her ad ya ham sütun adı ya da bir rename hedefi."""
    raw_meta = {
        "series_name", "study_date", "recon_kernel", "conv_kernel", "k_strength",
        "kvp", "mas", "fov", "slice_thickness_mm", "pixel_spacing_mm",
        "repeat", "group", "section", "label", "insert_size",
        "n_slices", "roi_mean_hu", "roi_std_hu", "voxel_count", "me_kev",
    }
    renamed = set(pipeline._COMPACT_RENAME.values())
    dropped = set(pipeline._COMPACT_DROP)

    for column in pipeline._COMPACT_META_ORDER + pipeline._COMPACT_TRAILING:
        assert column in renamed or column in raw_meta, (
            f"{column} ne ham sütun ne de bir rename hedefi — düzen bozulur"
        )

    # Rename kaynakları gerçekten var olan ham sütunlar olmalı.
    assert set(pipeline._COMPACT_RENAME) <= raw_meta
    # Bir sütun hem atılıp hem yeniden adlandırılamaz.
    assert not (dropped & set(pipeline._COMPACT_RENAME))


def test_compact_wide_produces_the_expected_layout(pipeline):
    """Uçtan uca: ham geniş tablo -> görünen sütun düzeni."""
    raw = pd.DataFrame({
        "series_name": ["USZ_A_Qr40f_120_100_300_1_PolyE_take1_recon1mm"],
        "study_date": ["20260402"], "recon_kernel": ["Qr40f"],
        "conv_kernel": ["Qr40"], "k_strength": [3], "kvp": [120.0],
        "mas": [100.0], "fov": [300.0], "slice_thickness_mm": [1.0],
        "pixel_spacing_mm": [0.586], "repeat": [1], "group": ["Inserts"],
        "section": ["INSERT1"], "label": ["Insert_01"], "insert_size": ["10mm"],
        "n_slices": [30], "roi_mean_hu": [520.0], "roi_std_hu": [7.0],
        "voxel_count": [5070], "me_kev": [np.nan],
        "stat_mean": [520.4], "cm_contrast_d1_3d_avg_fbs_w25.0": [12.3],
    })

    wide = pipeline._compact_wide(raw)
    columns = list(wide.columns)

    assert columns[:len(EXPECTED_META_BLOCK)] == EXPECTED_META_BLOCK
    assert columns[-len(EXPECTED_TRAILING):] == EXPECTED_TRAILING
    # Öznitelikler ortada kalmalı, meta bloğunu bölmemeli.
    assert set(columns[len(EXPECTED_META_BLOCK):-len(EXPECTED_TRAILING)]) == {
        "stat_mean", "cm_contrast_d1_3d_avg_fbs_w25.0",
    }
    # Atılanlar gerçekten gitmiş olmalı.
    for dropped in ("study_date", "conv_kernel", "roi_mean_hu", "n_slices"):
        assert dropped not in columns


def _wide_with_series() -> pd.DataFrame:
    return pd.DataFrame({
        "label": ["Insert_01", "Honeylarge"],
        "insert_size": ["10mm", ""],
        "series": ["USZ_A", "USZ_B"],
        "n_vox": [5070, 72030],
        "recon_k": ["Qr40f", "Qr40f"], "k_strength": [3, 3],
        "kVp": [120.0, 120.0], "mAs": [100.0, 100.0], "FOV": [300.0, 500.0],
        "SliceThickness": [1.0, 1.0], "me_keV": [None, None],
        "stat_mean": [520.4, -113.9],
        "section": ["INSERT1", "STAR16"],
    })


DICOM_BY_SERIES = {
    "USZ_A": {"dcm_StudyDate": "20260402", "dcm_Manufacturer": "Siemens",
              "dcm_StationName": "MT10102248", "dcm_KVP": "120"},
    "USZ_B": {"dcm_StudyDate": "20260402", "dcm_Manufacturer": "Siemens",
              "dcm_StationName": "MT10102248", "dcm_KVP": "120"},
}


def test_dicom_columns_are_attached_to_the_feature_sheet(pipeline):
    """Etiketler ayrı sayfada da duruyor ama satırı okurken elde olmalı."""
    wide = pipeline._attach_dicom_columns(_wide_with_series(), DICOM_BY_SERIES)
    columns = list(wide.columns)

    for name in ("dcm_StudyDate", "dcm_Manufacturer", "dcm_StationName"):
        assert name in columns
    assert wide["dcm_Manufacturer"].tolist() == ["Siemens", "Siemens"]


def test_attached_dicom_columns_sit_between_meta_and_features(pipeline):
    wide = pipeline._attach_dicom_columns(_wide_with_series(), DICOM_BY_SERIES)
    columns = list(wide.columns)

    assert columns.index("dcm_StudyDate") > columns.index("me_keV")
    assert columns.index("dcm_StationName") < columns.index("stat_mean")
    # Promoted olanlar önce, gerisi onları izler.
    assert columns.index("dcm_StudyDate") < columns.index("dcm_StationName")
    assert columns[-1] == "section"


def test_already_parsed_parameters_are_not_promoted_twice(pipeline):
    """kVp zaten kendi sütununda; dcm_KVP mavi bloğa alınıp iki kez gösterilmez."""
    assert "dcm_KVP" not in pipeline._DICOM_PROMOTED_COLUMNS
    assert "dcm_SliceThickness" not in pipeline._DICOM_PROMOTED_COLUMNS
    # Ama tabloda kaybolmuyor, sarı blokta duruyor.
    wide = pipeline._attach_dicom_columns(_wide_with_series(), DICOM_BY_SERIES)
    assert "dcm_KVP" in wide.columns


def test_attach_is_a_no_op_without_metadata(pipeline):
    """Metadata yoksa tablo aynen kalmalı — eksiklik çıktıyı bozmamalı."""
    original = _wide_with_series()
    assert list(pipeline._attach_dicom_columns(original, {}).columns) == \
        list(original.columns)


def test_block_colours_tint_the_meta_columns(pipeline, tmp_path):
    """Sol blok soruya göre renkli; öznitelik sütunları ROI bandını korumalı."""
    from openpyxl import Workbook

    wide = pd.DataFrame({
        "label": ["Insert_01"], "insert_size": ["10mm"],
        "series": ["USZ_A"], "n_vox": [5070],
        "recon_k": ["Qr40f"], "k_strength": [3], "kVp": [120.0],
        "mAs": [100.0], "FOV": [300.0], "SliceThickness": [1.0],
        "me_keV": [None], "stat_mean": [520.4], "section": ["INSERT1"],
    })

    wb = Workbook()
    ws = wb.active
    ws.append(list(wide.columns))
    ws.append(list(wide.iloc[0]))
    pipeline._apply_block_colors(ws, wide)

    def fill(column: str) -> str:
        return str(ws.cell(2, list(wide.columns).index(column) + 1)
                   .fill.start_color.rgb)[-6:]

    assert fill("label") == pipeline._BLOCK_FILL_ROI
    assert fill("insert_size") == pipeline._BLOCK_FILL_ROI
    assert fill("series") == pipeline._BLOCK_FILL_BOOKKEEPING
    assert fill("n_vox") == pipeline._BLOCK_FILL_BOOKKEEPING
    assert fill("kVp") == pipeline._BLOCK_FILL_ACQUISITION
    assert fill("me_keV") == pipeline._BLOCK_FILL_ACQUISITION
    # Öznitelik sütununa dokunulmamalı — orada satır bandı okunuyor.
    assert fill("stat_mean") not in {
        pipeline._BLOCK_FILL_ROI, pipeline._BLOCK_FILL_BOOKKEEPING,
        pipeline._BLOCK_FILL_ACQUISITION,
    }


# --------------------------------------------------------------------------- #
# Sütun genişlikleri — seri adı kırpılamaz
# --------------------------------------------------------------------------- #

def _styled_feature_sheet(pipeline):
    """Gerçek yoldan üretilmiş bir feature sayfası: _compact_wide -> _style_sheet."""
    from openpyxl import Workbook

    raw = pd.DataFrame({
        "series_name": [
            "Radiologie_USZ_Siemens_Healthineers_NAEOTOM_Alpha_20260402_"
            "Qr40f_Q3_120_100_300_1_PolyE_take1_recon1mm",
        ],
        "study_date": ["20260402"], "recon_kernel": ["Qr40f"],
        "conv_kernel": ["Qr40f"], "k_strength": [3], "kvp": [120.0],
        "mas": [100.0], "fov": [300.0], "slice_thickness_mm": [1.0],
        "pixel_spacing_mm": [0.586], "repeat": [""], "group": ["Inserts"],
        "section": ["5mm"], "label": ["Insert_01"], "insert_size": ["5mm"],
        "n_slices": [30], "roi_mean_hu": [520.0], "roi_std_hu": [7.0],
        "voxel_count": [1470], "me_kev": [np.nan], "stat_mean": [520.4],
    })
    wide = pipeline._compact_wide(raw)

    wb = Workbook()
    ws = wb.active
    ws.append(list(wide.columns))
    ws.append(list(wide.iloc[0]))
    pipeline._style_sheet(ws, wide,
                         freeze_meta_count=pipeline._sheet_meta_count(wide))
    return ws, wide


def _width(ws, wide, column: str) -> float:
    from openpyxl.utils import get_column_letter

    letter = get_column_letter(list(wide.columns).index(column) + 1)
    return ws.column_dimensions[letter].width


def test_series_column_is_not_truncated_wherever_it_sits(pipeline):
    """Seri adı 60 birim almalı — A sütununda olmasa bile.

    Bu muafiyet `col_idx == 1`e sabitlenmişti. Meta bloğu yeniden gruplanınca
    A sütununu `label` (9 karakter) devraldı ve 60'ı miras aldı; `series`
    (43 karakter) C sütununa kayıp 14 karakter sınırına düştü — muafiyetin
    önlemek için var olduğu çakışmanın tam kendisi. Artık başlık adına bağlı.
    """
    ws, wide = _styled_feature_sheet(pipeline)

    assert list(wide.columns)[0] == "label", "meta düzeni değişti, test güncellenmeli"
    assert _width(ws, wide, "series") == pipeline._WIDE_DISPLAY_WIDTH
    # Kısa tanımlayıcılar geniş muafiyeti almamalı, yoksa ekranı boşa harcar.
    assert _width(ws, wide, "label") < pipeline._WIDE_DISPLAY_WIDTH
    assert _width(ws, wide, "insert_size") < pipeline._WIDE_DISPLAY_WIDTH


def test_wide_exemption_is_matched_by_name_not_position(pipeline):
    """Muafiyet konuma değil ada bakmalı."""
    assert "series" in pipeline._WIDE_DISPLAY_COLUMNS
    # Sabit gerçekten kullanılıyor mu — çağrı yeri sayı gömmüş olmasın.
    source = (PROJECT_DIR /
              "extract_pcct_phantom_radiomics_dicom_inserts_only_v4.py"
              ).read_text(encoding="utf-8")
    assert "_WIDE_DISPLAY_COLUMNS" in source
    assert "col_idx == 1" not in source, "genişlik yine konuma sabitlenmiş"


def test_series_sheet_keeps_the_wide_series_column(pipeline):
    """`series` sayfasında seri adı hâlâ 1. sütun; muafiyet orada da geçerli."""
    from openpyxl import Workbook

    disp = pd.DataFrame({
        "series": ["Qr40f_Q3_120_100_300_1_PolyE_take1_recon1mm"],
        "date": ["20260402"], "recon_k": ["Qr40f"], "k_strength": [3],
        "kVp": [120.0], "mAs": [100.0], "FOV": [300.0],
        "thick_mm": [1.0], "pixel_mm": [0.586], "repeat": [""],
    })
    wb = Workbook()
    ws = wb.active
    ws.append(list(disp.columns))
    ws.append(list(disp.iloc[0]))
    pipeline._style_sheet(ws, disp, freeze_meta_count=1, as_table=True,
                         table_name="SeriesSummary")

    assert _width(ws, disp, "series") == pipeline._WIDE_DISPLAY_WIDTH


def test_ranking_sheet_does_not_waste_a_60_wide_rank_column(pipeline):
    """insert_ranking'in A sütunu `rank`; 60 birim genişlik anlamsızdı."""
    from openpyxl import Workbook

    rank = pd.DataFrame({"rank": [1], "label": ["Insert_07"],
                         "mean_HU_global": [-98.4]})
    wb = Workbook()
    ws = wb.active
    ws.append(list(rank.columns))
    ws.append(list(rank.iloc[0]))
    pipeline._style_sheet(ws, rank, freeze_meta_count=2)

    assert _width(ws, rank, "rank") < pipeline._WIDE_DISPLAY_WIDTH


def test_the_two_workbooks_use_the_same_colour_scheme(pipeline):
    """İki hattın çıktısı aynı şekilde okunmalı: yeşil=ne, sarı=kayıt, mavi=koşul.

    Sabitler kaynaktan okunuyor, modül yüklenmiyor: phantomxpro'nun
    excel_report'u kendi `config`'ine bağlı ve buradan çalıştırılamıyor.
    """
    import re

    path = organ_radiomics_project() / "excel_report.py"
    if not path.is_file():
        pytest.skip(f"{path} yok")
    source = path.read_text(encoding="utf-8")

    def constant(name: str) -> str:
        match = re.search(rf'^{name}\s*=\s*"([0-9A-Fa-f]{{6}})"', source, re.MULTILINE)
        assert match, f"{name} phantomxpro excel_report.py içinde bulunamadı"
        return match.group(1)

    assert pipeline._BLOCK_FILL_ROI == constant("_ORGAN_FILL")
    assert pipeline._BLOCK_FILL_BOOKKEEPING == constant("_DICOM_FILL")
    assert pipeline._BLOCK_FILL_ACQUISITION == constant("_ACQUISITION_FILL")


# --------------------------------------------------------------------------- #
# Beyan sayfası
# --------------------------------------------------------------------------- #

def test_methods_rows_cover_every_reportable_item(pipeline):
    rows = pipeline._methods_declaration_rows(pd.DataFrame(), n_series=3)
    sections = {section for section, _, _ in rows}
    assert {"Dataset", "Software", "Image", "ROI", "Re-segmentation",
            "Discretisation", "Texture", "Features", "Caveats"} <= sections

    items = {item: value for _, item, value in rows}
    assert items["Bin width (HU)"] == "25"
    assert items["Grey-level origin (HU)"] == "-2000"
    assert "NONE" in items["Interpolation / resampling"]
    assert "cuboid" in items["Analysis mask shape"]


def test_methods_rows_are_derived_not_hardcoded(pipeline, monkeypatch):
    """Beyan, sabitler değişince onunla birlikte değişmeli — yoksa eskir."""
    monkeypatch.setattr(pipeline, "_DISCRETISATION_ANCHOR_HU", -1500.0)
    monkeypatch.setattr(pipeline, "DISCRETISATION_RANGE_HU", (-1500.0, 31743.0))

    rows = pipeline._methods_declaration_rows(pd.DataFrame())
    items = {item: value for _, item, value in rows}
    assert items["Grey-level origin (HU)"] == "-1500"
    assert items["Range (HU)"] == "(-1500, 31743)"


def test_methods_sheet_is_written_first(pipeline, tmp_path):
    """Kitap açıldığında önce nasıl üretildiği görünmeli."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "3D_base"
    pipeline._write_methods_sheet(wb, pd.DataFrame(), n_series=2,
                                  dataset_root=str(tmp_path))

    assert wb.sheetnames[0] == "methods"
    ws = wb["methods"]
    assert ws.freeze_panes is not None
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Ready to paste" in text
    assert "fixed bin size of 25 HU" in text


def test_no_table_of_contents_sheet(pipeline):
    """TOC kaldırıldı: gezinme değeri yoktu, beyan sayfası onun yerini aldı."""
    assert not hasattr(pipeline, "_write_toc_sheet")
    assert "TOC" not in pipeline._TAB_COLORS
    assert "methods" in pipeline._TAB_COLORS
